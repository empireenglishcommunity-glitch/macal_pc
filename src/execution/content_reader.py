"""
File Content Reader - Extracts text from various file types.
Enables the agent to read, search, and summarize file contents.
"""
from pathlib import Path


def read_file_content(filepath, max_chars=5000):
    """Read text content from a file. Returns first max_chars characters."""
    p = Path(filepath).expanduser()
    if not p.exists():
        return {"success": False, "error": "File not found: " + str(p)}
    if not p.is_file():
        return {"success": False, "error": "Not a file: " + str(p)}

    ext = p.suffix.lower()
    try:
        if ext in (".txt", ".md", ".py", ".js", ".json", ".csv", ".html",
                   ".css", ".xml", ".log", ".bat", ".sh", ".yaml", ".yml",
                   ".toml", ".ini", ".cfg"):
            text = p.read_text(encoding="utf-8", errors="replace")
        elif ext == ".pdf":
            text = _read_pdf(p)
        elif ext == ".docx":
            text = _read_docx(p)
        elif ext in (".xlsx", ".xls"):
            text = _read_excel(p)
        else:
            text = "[Binary or unsupported file type: " + ext + "]"

        if len(text) > max_chars:
            text = text[:max_chars] + "\n... (truncated at " + str(max_chars) + " chars)"

        return {"success": True, "content": text, "chars": len(text), "type": ext}
    except Exception as e:
        return {"success": False, "error": str(e)}


def _read_pdf(filepath):
    """Extract text from PDF using pymupdf."""
    import pymupdf
    doc = pymupdf.open(str(filepath))
    text = ""
    for page in doc:
        text += page.get_text()
    doc.close()
    return text


def _read_docx(filepath):
    """Extract text from Word document."""
    from docx import Document
    doc = Document(str(filepath))
    text = ""
    for para in doc.paragraphs:
        text += para.text + "\n"
    return text


def _read_excel(filepath):
    """Read Excel file as text summary."""
    from openpyxl import load_workbook
    wb = load_workbook(str(filepath), read_only=True, data_only=True)
    text = ""
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        text += "Sheet: " + sheet_name + "\n"
        row_count = 0
        for row in ws.iter_rows(values_only=True):
            if row_count > 50:
                text += "... (more rows)\n"
                break
            text += " | ".join(str(cell) if cell is not None else "" for cell in row) + "\n"
            row_count += 1
    wb.close()
    return text


def search_in_file(filepath, query):
    """Search for a text string inside a file. Returns matching lines."""
    result = read_file_content(filepath, max_chars=50000)
    if not result["success"]:
        return result

    lines = result["content"].split("\n")
    matches = []
    query_lower = query.lower()
    for i, line in enumerate(lines, 1):
        if query_lower in line.lower():
            matches.append({"line": i, "text": line.strip()})

    return {"success": True, "query": query, "matches": matches, "total_matches": len(matches)}
